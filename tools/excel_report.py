from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXCEL_CONFIG
import os
from datetime import datetime


class ExcelReport:

    def __init__(self, output_dir):
        self.output_dir = output_dir
        self.date = datetime.now().strftime('%Y%m%d')
        self.styles = self._create_styles()

        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

    def _create_styles(self):
        return {
            'header_font': Font(**EXCEL_CONFIG['header_font']),
            'header_fill': PatternFill(**EXCEL_CONFIG['header_fill']),
            'header_alignment': Alignment(horizontal='center', vertical='center', wrap_text=True),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            ),
            'row_height': EXCEL_CONFIG['row_height'],
            'col_width': EXCEL_CONFIG['col_width']
        }

    def _set_column_width(self, worksheet, column_mapping):
        for col_name, width in column_mapping.items():
            col_letter = get_column_letter(col_name)
            worksheet.column_dimensions[col_letter].width = width

    def _set_header_style(self, worksheet):
        for col in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=1, column=col)
            cell.font = self.styles['header_font']
            cell.fill = self.styles['header_fill']
            cell.alignment = self.styles['header_alignment']
            cell.border = self.styles['border']

    def _set_cell_style(self, worksheet, row, col):
        cell = worksheet.cell(row=row, column=col)
        cell.border = self.styles['border']
        cell.alignment = Alignment(horizontal='center', vertical='center')

    def _add_summary_info(self, worksheet, data_dict, start_row=2):
        for idx, (key, info) in enumerate(data_dict.items()):
            row = start_row + idx
            worksheet.cell(row=row, column=1, value=info['label'])
            worksheet.cell(row=row, column=2, value=info['value'])

    def _format_percent(self, value):
        if value is not None:
            return f"{value:.2f}%"
        return "-"

    def generate_report(self, report_name, data, headers, sheet_name=None):
        """
        生成报表
        """
        if not sheet_name:
            sheet_name = report_name

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.row_dimensions.height = self.styles['row_height']
        worksheet.append(headers)

        self._set_header_style(worksheet)

        column_width = self.styles['col_width']
        if len(headers) > 0:
            self._set_column_width(worksheet, {
                headers[0]: column_width.get('name', column_width['default']),
                headers[1]: column_width.get('number', column_width['default']),
                headers[2]: column_width['default']
            })

        for row_idx, row_data in enumerate(data, start=2):
            for col_idx, value in enumerate(row_data, start=1):
                cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
                self._set_cell_style(worksheet, row_idx, col_idx)

        filename = f"{report_name}_{self.date}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        workbook.save(filepath)
        print(f"The report has been generated: {filepath}")

        return filepath

    def generate_summary_report(self, report_name, data_dict):
        """
        生成汇总报表（Key-Value）
        """
        sheet_name = f"{report_name}_Summary"

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.row_dimensions.height = self.styles['row_height']

        headers = ['指标', '数值']
        worksheet.append(headers)

        self._set_header_style(worksheet)

        self._set_column_width(worksheet, {
            '指标': 30,
            '数值': 30
        })

        for idx, (key, value) in enumerate(data_dict.items(), start=2):
            worksheet.cell(row=idx, column=1, value=key)
            worksheet.cell(row=idx, column=2, value=value)
            self._set_cell_style(worksheet, idx, 1)
            self._set_cell_style(worksheet, idx, 2)

        filename = f"{report_name}_{self.date}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        workbook.save(filepath)
        print(f"The summary report has been generated: {filepath}")

        return filepath

    def generate_comparison_report(self, report_name, group1_data, group2_data):
        """
        生成对比报表
        """
        sheet_name = f"{report_name}_Comparison"

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.row_dimensions.height = self.styles['row_height']

        headers = ['项目', group1_name, group2_name, '差异']
        worksheet.append(headers)

        self._set_header_style(worksheet)

        self._set_column_width(worksheet, {
            '项目': 25,
            group1_name: 20,
            group2_name: 20,
            '差异': 20
        })

        for idx, (row1, row2) in enumerate(zip(group1_data, group2_data), start=2):
            worksheet.cell(row=idx, column=1, value=row1[0])
            worksheet.cell(row=idx, column=2, value=row1[1])

            worksheet.cell(row=idx, column=3, value=row2[1])

            val1 = float(row1[1]) if row1[1] else 0
            val2 = float(row2[1]) if row2[1] else 0
            diff = val2 - val1
            worksheet.cell(row=idx, column=4, value=diff)

            for col in range(1, 5):
                self._set_cell_style(worksheet, idx, col)

        filename = f"{report_name}_{self.date}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        workbook.save(filepath)
        print(f"The compared report has been generated: {filepath}")

        return filepath
